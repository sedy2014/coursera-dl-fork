"""Render validated JSON objects using trusted, lazily imported PPTX helpers.

No source strings are parsed or executed. Visible whitespace is collapsed for
layout; speaker notes are retained on every continuation of their source slide,
with CRLF and CR line endings normalized to LF.
The aggregate input text limit is 200,000 Unicode characters.
"""

import os
from pathlib import Path
import tempfile
import unicodedata


MAX_TOTAL_CHARS = 200_000
PALETTES = ("IBM New", "GCP", "IBM Original")
BODY_TOP = 160
BODY_BOTTOM = 486
BODY_LINE_HEIGHT = 22
BULLET_GAP = 8


def _check_keys(value, keys, location):
    if type(value) is not dict or value.keys() != keys:
        raise ValueError(f"{location} must contain exactly {sorted(keys)}")


def _check_text(value, limit, location, nonempty=False):
    if type(value) is not str or len(value) > limit:
        raise ValueError(f"{location} must be a string of at most {limit} characters")
    if nonempty and not value.strip():
        raise ValueError(f"{location} must not be empty")
    if any(
        ord(char) < 32 and char not in "\t\n\r"
        or 0xD800 <= ord(char) <= 0xDFFF
        or ord(char) in (0xFFFE, 0xFFFF)
        for char in value
    ):
        raise ValueError(f"{location} contains characters unsupported by XML")
    return value


def validate_deck(data) -> dict:
    """Return a detached validated JSON object, or raise ValueError.

    Only built-in dict/list/str values are accepted, not JSON or Python source
    strings. Empty slide titles and notes are allowed. All keys are required.
    """
    _check_keys(data, {"title", "slides"}, "deck")
    title = _check_text(data["title"], 120, "deck.title", nonempty=True)
    slides = data["slides"]
    if type(slides) is not list or not 1 <= len(slides) <= 60:
        raise ValueError("deck.slides must be a list of 1..60 slides")
    total = len(title)
    validated = []
    for index, slide in enumerate(slides):
        location = f"slides[{index}]"
        _check_keys(slide, {"title", "bullets", "notes"}, location)
        slide_title = _check_text(slide["title"], 100, f"{location}.title")
        notes = _check_text(slide["notes"], 8000, f"{location}.notes")
        bullets = slide["bullets"]
        if type(bullets) is not list or not 1 <= len(bullets) <= 6:
            raise ValueError(f"{location}.bullets must be a list of 1..6 strings")
        bullets = [
            _check_text(bullet, 400, f"{location}.bullets[{number}]", nonempty=True)
            for number, bullet in enumerate(bullets)
        ]
        total += len(slide_title) + len(notes) + sum(map(len, bullets))
        if total > MAX_TOTAL_CHARS:
            raise ValueError(f"deck text exceeds {MAX_TOTAL_CHARS} characters")
        validated.append({"title": slide_title, "bullets": bullets, "notes": notes})
    return {"title": title, "slides": validated}


def _units(text):
    return sum(2 if unicodedata.east_asian_width(char) in "WF" else 1 for char in text)


def _wrap_text(text, width, max_lines=None, prefer_words=True):
    """Wrap at word boundaries, splitting long tokens with wide glyph allowance."""
    lines = []
    line = ""
    for char in " ".join(text.split()):
        if _units(line + char) > width:
            before, separator, after = line.rpartition(" ")
            if char == " ":
                lines.append(line)
                line = ""
                continue
            if separator and prefer_words:
                lines.append(before)
                line = after
            else:
                lines.append(line)
                line = ""
        line += char
    if line:
        lines.append(line)
    if max_lines is not None and len(lines) > max_lines:
        return _wrap_text(text, width, prefer_words=False)
    return lines or [""]


def _paginate(bullets):
    pages = []
    page = []
    used = 0
    available = BODY_BOTTOM - BODY_TOP
    for bullet in bullets:
        lines = _wrap_text(bullet, 86)
        height = len(lines) * BODY_LINE_HEIGHT + BULLET_GAP
        if page and used + height > available:
            pages.append(page)
            page = []
            used = 0
        while lines:
            capacity = (available - used - BULLET_GAP) // BODY_LINE_HEIGHT
            chunk, lines = lines[:capacity], lines[capacity:]
            page.append(chunk)
            used += len(chunk) * BODY_LINE_HEIGHT + BULLET_GAP
            if lines:
                pages.append(page)
                page = []
                used = 0
    pages.append(page)
    return pages


def _textbox(helpers, slide, name, left, top, width, height, lines,
             size, spacing, color, bold=False):
    from pptx.enum.text import MSO_AUTO_SIZE

    shape = helpers.txbox(
        slide, helpers.Pt(left), helpers.Pt(top), helpers.Pt(width),
        helpers.Pt(height), "", size=size, color=color, bold=bold, wrap=False,
    )
    shape.name = name
    frame = shape.text_frame
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.text = "\v".join(lines)
    for paragraph in frame.paragraphs:
        paragraph.space_before = helpers.Pt(0)
        paragraph.space_after = helpers.Pt(0)
        paragraph.line_spacing = helpers.Pt(spacing)
        for run in paragraph.runs:
            run.font.name = "Courier New"
            run.font.size = helpers.Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return shape


def _validate_saved(path, expected_notes):
    from pptx import Presentation

    reopened = Presentation(str(path))
    if not reopened.slides or len(reopened.slides) != len(expected_notes):
        raise ValueError("Saved presentation has an invalid slide count")
    for slide, notes in zip(reopened.slides, expected_notes):
        if not slide.has_notes_slide:
            raise ValueError("Saved presentation is missing speaker notes")
        frame = slide.notes_slide.notes_text_frame
        if frame is None or frame.text != notes:
            raise ValueError("Saved presentation speaker notes did not round-trip")


def render_deck(data, output_path, palette="IBM New") -> str:
    """Atomically save a validated deck and return its resolved output path.

    The output parent must exist. Input slides may expand into continuation
    slides. Palette is application-controlled, never a field in generated JSON.
    Validation and importing this module do not require python-pptx.
    """
    deck = validate_deck(data)
    if type(palette) is not str or palette not in PALETTES:
        raise ValueError(f"palette must be one of {PALETTES}")

    from . import slide_utils as helpers

    colors = {
        "IBM New": (helpers.IBMBG, helpers.IBMBLUE, helpers.NAVY),
        "GCP": (helpers.GCPSLATE, helpers.GCPBLUE, helpers.CHARCOAL),
        "IBM Original": (helpers.ORIGSLATE, helpers.ORIGBLUE, helpers.ORIGNAVY),
    }
    background, accent, foreground = colors[palette]
    presentation, blank = helpers.new_presentation()
    presentation.core_properties.title = deck["title"]
    width = presentation.slide_width.pt
    expected_notes = []
    for source in deck["slides"]:
        for continuation, page in enumerate(_paginate(source["bullets"]), start=1):
            slide = presentation.slides.add_slide(blank)
            helpers.bg(slide, background)
            helpers.rect(slide, 0, 0, helpers.W, helpers.Pt(5), accent)
            title = source["title"]
            if continuation > 1:
                title += f" (continued {continuation})"
            _textbox(helpers, slide, "Title", 40, 20, width - 80, 120,
                     _wrap_text(title, 64, max_lines=4), 22, 28, foreground, bold=True)
            top = BODY_TOP
            for number, lines in enumerate(page, start=1):
                height = len(lines) * BODY_LINE_HEIGHT
                display = ["\u2022 " + lines[0]] + ["  " + line for line in lines[1:]]
                _textbox(helpers, slide, f"Bullet {number}", 40, top,
                         width - 80, height, display, 16, BODY_LINE_HEIGHT, foreground)
                top += height + BULLET_GAP
            _textbox(helpers, slide, "Footer", 40, 502, width - 140, 28,
                     _wrap_text(deck["title"], 128, max_lines=2), 9, 12, foreground)
            _textbox(helpers, slide, "Page number", width - 90, 502, 60, 28,
                     [str(len(expected_notes) + 1)], 9, 12, accent)
            notes = source["notes"].replace("\r\n", "\n").replace("\r", "\n")
            helpers.set_notes(slide, notes)
            expected_notes.append(notes)

    destination = Path(output_path).expanduser().resolve()
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=destination.parent, prefix=".slide-render-", suffix=".pptx", delete=False,
        ) as handle:
            temporary = Path(handle.name)
        presentation.save(str(temporary))
        _validate_saved(temporary, expected_notes)
        os.replace(temporary, destination)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return str(destination)
