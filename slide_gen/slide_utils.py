"""
slide_gen/slide_utils.py
------------------------
Shared python-pptx helpers and brand palettes for LLM-generated slide scripts.

Generated scripts start with:
    from slide_gen.slide_utils import *
    prs, blank = new_presentation()
    s1 = prs.slides.add_slide(blank)
    bg(s1, IBMBG)
    title_bar(s1, "My Title", accent=IBMBLUE)
    dual_tags(s1, day="Monday")
    prs.save(output_path)

Layout constants
----------------
W  = Inches(13.33)   slide width
H  = Inches(7.5)     slide height
CT = Pt(82)          content top (below title bar)
GAP = Pt(12)         default gap between elements

Standard three-column geometry (for card grids):
    cw  = (W - Pt(24)*2 - GAP*2) / 3
    L0  = Pt(24)
    L1  = L0 + cw + GAP
    L2  = L1 + cw + GAP
    ch  = H - CT - Pt(8)

Two-column geometry:
    hw = (W - Pt(24)*2 - GAP) / 2
    rx = Pt(24) + hw + GAP
    mh = H - CT - Pt(8)

Horizontal pipeline (N stages):
    sw = (W - Pt(24)*2 - (N-1)*GAP) / N
    for i in range(N):
        lx = Pt(24) + i*(sw+GAP)
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ---------------------------------------------------------------------------
# Slide dimensions
# ---------------------------------------------------------------------------

W   = Inches(13.33)
H   = Inches(7.5)
CT  = Pt(82)      # content top — below title bar
GAP = Pt(12)      # default gap between elements

# ---------------------------------------------------------------------------
# Shared neutrals (present in all three palette families)
# ---------------------------------------------------------------------------

WHITE    = RGBColor(0xff, 0xff, 0xff)
NAVY     = RGBColor(0x08, 0x18, 0x30)   # IBM-new near-black
CHARCOAL = RGBColor(0x21, 0x22, 0x26)   # GCP near-black
GRAY     = RGBColor(0x52, 0x56, 0x5b)
LGRAY    = RGBColor(0xc1, 0xc7, 0xcd)
DARK_BLUE = RGBColor(0x00, 0x1d, 0x6c)

# ---------------------------------------------------------------------------
# IBM New palette  (use for IBM Data Science Professional Certificate)
# ---------------------------------------------------------------------------

IBMBLUE   = RGBColor(0x0f, 0x62, 0xfe)
IBMCYAN   = RGBColor(0x11, 0x92, 0xe8)
IBMGREEN  = RGBColor(0x19, 0x80, 0x38)
IBMPURPLE = RGBColor(0x7c, 0x3d, 0xd6)
IBMTEAL   = RGBColor(0x00, 0x9d, 0x9a)
IBMRED    = RGBColor(0xda, 0x1e, 0x28)
IBMYELLOW = RGBColor(0xf1, 0xc2, 0x1b)
IBMBG     = RGBColor(0xf2, 0xf4, 0xf8)   # cool gray 10 — default slide background

LT_IBMBLUE   = RGBColor(0xe8, 0xed, 0xff)
LT_IBMCYAN   = RGBColor(0xe5, 0xf6, 0xff)
LT_IBMGREEN  = RGBColor(0xde, 0xf7, 0xe9)
LT_IBMPURPLE = RGBColor(0xee, 0xe8, 0xff)
LT_IBMTEAL   = RGBColor(0xd9, 0xfb, 0xfb)
LT_IBMRED    = RGBColor(0xfd, 0xe8, 0xe8)
LT_IBMYELLOW = RGBColor(0xfe, 0xf9, 0xe0)

# ---------------------------------------------------------------------------
# GCP palette  (use for Google Cloud / Vertex AI courses)
# ---------------------------------------------------------------------------

GCPBLUE   = RGBColor(0x1a, 0x73, 0xe8)
GCPGREEN  = RGBColor(0x18, 0x97, 0x49)
GCPYELLOW = RGBColor(0xf2, 0x99, 0x00)
GCPRED    = RGBColor(0xea, 0x43, 0x35)
GCPPURPLE = RGBColor(0x7b, 0x2f, 0xbe)
GCPTEAL   = RGBColor(0x00, 0x89, 0x7b)
GCPINDIGO = RGBColor(0x3d, 0x50, 0xa1)
GCPSLATE  = RGBColor(0xf1, 0xf3, 0xf4)  # GCP background

LT_GCPBLUE   = RGBColor(0xe8, 0xf0, 0xfe)
LT_GCPGREEN  = RGBColor(0xe6, 0xf4, 0xea)
LT_GCPYELLOW = RGBColor(0xfe, 0xf7, 0xe0)
LT_GCPRED    = RGBColor(0xfc, 0x8e, 0x8e)  # approx
LT_GCPPURPLE = RGBColor(0xf3, 0xe8, 0xfd)
LT_GCPTEAL   = RGBColor(0xe6, 0xf4, 0xf1)

# ---------------------------------------------------------------------------
# IBM Original palette  (use for make_ppt1/2/3 style — "What is Data Science" era)
# ---------------------------------------------------------------------------

ORIGBLUE   = RGBColor(0x1a, 0x56, 0xdb)
ORIGGREEN  = RGBColor(0x0a, 0x8a, 0x5a)
ORIGORANGE = RGBColor(0xe0, 0x70, 0x00)
ORIGPURPLE = RGBColor(0x6b, 0x21, 0xa8)
ORIGRED    = RGBColor(0xdc, 0x26, 0x26)
ORIGSLATE  = RGBColor(0xf0, 0xf2, 0xf8)   # original IBM background
ORIGNAVY   = RGBColor(0x0d, 0x1b, 0x4b)

LT_ORIGBLUE   = RGBColor(0xdb, 0xe8, 0xfb)
LT_ORIGGREEN  = RGBColor(0xd5, 0xf0, 0xe4)
LT_ORIGORANGE = RGBColor(0xfd, 0xf0, 0xe0)
LT_ORIGPURPLE = RGBColor(0xef, 0xe0, 0xff)

# Tag colors
DELEGATE_COL = IBMGREEN
PROF_COL     = IBMBLUE

# ---------------------------------------------------------------------------
# Presentation factory
# ---------------------------------------------------------------------------


def new_presentation():
    """Return (prs, blank_layout) ready for a 13.33 × 7.5 inch slide deck."""
    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    return prs, blank


# ---------------------------------------------------------------------------
# Primitive helpers
# ---------------------------------------------------------------------------


def bg(sl, color=IBMBG):
    """Fill the entire slide with *color*."""
    sh = sl.shapes.add_shape(1, 0, 0, W, H)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()


def rect(sl, l, t, w, h, fill,
         line=None, lw=Pt(0.75),
         rounded=False, radius=0.04):
    """
    Draw a filled rectangle.

    Parameters
    ----------
    rounded : bool   Use rounded corners (shape type 5).
    radius  : float  Corner radius as a fraction of the shorter dimension.
    """
    sh = sl.shapes.add_shape(5 if rounded else 1, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = lw
    else:
        sh.line.fill.background()
    if rounded:
        sh.adjustments[0] = radius
    return sh


def txbox(sl, l, t, w, h, text,
          size=13, bold=False, color=NAVY,
          align=PP_ALIGN.LEFT, wrap=True, italic=False):
    """Add a text box. Returns the shape."""
    tb = sl.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    p   = tf.paragraphs[0]
    run = p.add_run()
    run.text            = text
    run.font.size       = Pt(size)
    run.font.bold       = bold
    run.font.italic     = italic
    run.font.color.rgb  = color
    p.alignment         = align
    return tb


def set_notes(sl, text):
    """Write *text* to the slide's speaker-notes pane."""
    sl.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------------------
# Layout components
# ---------------------------------------------------------------------------


def title_bar(sl, title, subtitle=None, accent=IBMBLUE):
    """
    IBM-new style title bar.
    Thin accent stripe at top, then 27pt title, optional italic subtitle.
    Leaves ~3.4 inches on the right for teacher-tag chips.
    """
    rect(sl, 0, 0, W, Pt(5), accent)
    txbox(sl, Pt(28), Pt(12), W - Inches(3.4), Pt(44),
          title, size=27, bold=True, color=NAVY)
    if subtitle:
        txbox(sl, Pt(28), Pt(56), W - Inches(3.4), Pt(22),
              subtitle, size=12, italic=True, color=GRAY)


def title_row(sl, title, subtitle=None, accent=GCPBLUE):
    """
    GCP style title bar.
    Same as title_bar but the title spans nearly the full width.
    """
    rect(sl, 0, 0, W, Pt(5), accent)
    txbox(sl, Pt(28), Pt(12), W - Pt(56), Pt(44),
          title, size=26, bold=True, color=NAVY)
    if subtitle:
        txbox(sl, Pt(28), Pt(56), W - Pt(56), Pt(22),
              subtitle, size=12, italic=True, color=GRAY)


def bottom_bar(sl, label, text, accent=IBMBLUE, t=None):
    """
    Colored summary bar that fills the slide from *t* to the bottom.
    *label* is a short uppercase eyebrow; *text* is the body.
    """
    if t is None:
        t = H - Pt(70)
    bh = H - t - Pt(4)
    rect(sl, 0, t, W, bh, accent)
    rect(sl, 0, t, Pt(6), bh, DARK_BLUE)
    txbox(sl, Pt(16), t + Pt(4), Inches(2.0), Pt(18),
          label, size=9, bold=True, color=WHITE)
    txbox(sl, Pt(16), t + Pt(20), W - Pt(32), bh - Pt(24),
          text, size=12, bold=False, color=WHITE, wrap=True)


def card(sl, l, t, w, h, label, title, bullets,
         accent=IBMBLUE, bg_color=WHITE):
    """
    Titled card with an accent top strip, an eyebrow label, a title line,
    and a bullet list.

    Parameters
    ----------
    bullets : list[str]   Each string becomes one bullet paragraph.
    """
    rect(sl, l, t, w, h, bg_color, accent, Pt(1.5), rounded=True, radius=0.03)
    rect(sl, l, t, w, Pt(5), accent)
    txbox(sl, l + Pt(10), t + Pt(8), w - Pt(18), Pt(16),
          label, size=9, bold=True, color=accent)
    txbox(sl, l + Pt(10), t + Pt(24), w - Pt(18), Pt(24),
          title, size=13, bold=True, color=NAVY)
    bx = sl.shapes.add_textbox(l + Pt(12), t + Pt(50), w - Pt(22), h - Pt(58))
    tf = bx.text_frame
    tf.word_wrap = True
    first = True
    for item in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_before = Pt(3)
        run = p.add_run()
        run.text           = item
        run.font.size      = Pt(12)
        run.font.color.rgb = CHARCOAL


def mbullets(sl, l, t, w, h, items, size=12, color=NAVY):
    """
    Add a vertical list of bullet strings to a text box.

    Parameters
    ----------
    items : list[str | tuple[str, bool]]
        Plain strings, or (text, bold) tuples.
    """
    bx = sl.shapes.add_textbox(l, t, w, h)
    tf = bx.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        if isinstance(item, tuple):
            txt, bld = item
        else:
            txt, bld = item, False
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_before = Pt(4)
        run = p.add_run()
        run.text           = txt
        run.font.size      = Pt(size)
        run.font.bold      = bld
        run.font.color.rgb = color


def code_box(sl, l, t, w, h, lines):
    """
    Dark monospace code box with simple syntax colouring.

    Parameters
    ----------
    lines : list[str]   Lines of code (no newlines needed).
                        Lines starting with '#' are rendered in green.
                        Lines starting with '>' are rendered in cyan.
    """
    rect(sl, l, t, w, h, CHARCOAL, rounded=True, radius=0.02)
    bx = sl.shapes.add_textbox(l + Pt(10), t + Pt(8), w - Pt(18), h - Pt(14))
    tf = bx.text_frame
    tf.word_wrap = False
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        run = p.add_run()
        run.text      = line
        run.font.size = Pt(9.5)
        run.font.name = "Courier New"
        if line.strip().startswith('#'):
            run.font.color.rgb = IBMGREEN
        elif line.strip().startswith('>'):
            run.font.color.rgb = IBMCYAN
        else:
            run.font.color.rgb = RGBColor(0xd4, 0xd4, 0xd4)


# ---------------------------------------------------------------------------
# Teacher-tag chips (top-right corner)
# ---------------------------------------------------------------------------


def dual_tags(sl, day="Monday"):
    """
    Two stacked chips: green DELEGATE chip + blue PROFESSOR chip.
    Used when the delegate teaches a Combined class and the professor
    must separately cover the same slides for Section A1.
    """
    tw = Inches(3.0)
    th = Pt(20)
    rect(sl, W - tw - Pt(8), Pt(4), tw, th, IBMGREEN, rounded=True, radius=0.06)
    txbox(sl, W - tw - Pt(8), Pt(4), tw, th,
          f"DELEGATE — {day} | Combined class",
          size=8.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    rect(sl, W - tw - Pt(8), Pt(27), tw, th, IBMBLUE, rounded=True, radius=0.06)
    txbox(sl, W - tw - Pt(8), Pt(27), tw, th,
          "PROFESSOR — also taught for Section A1",
          size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)


def prof_tag(sl, label="PROFESSOR — All sections"):
    """Single blue chip for professor-only slides."""
    tw = Inches(2.8)
    th = Pt(20)
    rect(sl, W - tw - Pt(8), Pt(4), tw, th, IBMBLUE, rounded=True, radius=0.06)
    txbox(sl, W - tw - Pt(8), Pt(4), tw, th,
          label, size=8.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)


def delegate_tag(sl, label="DELEGATE — Combined class"):
    """Single green chip for delegate-only slides."""
    tw = Inches(2.8)
    th = Pt(20)
    rect(sl, W - tw - Pt(8), Pt(4), tw, th, IBMGREEN, rounded=True, radius=0.06)
    txbox(sl, W - tw - Pt(8), Pt(4), tw, th,
          label, size=8.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)


def gcp_teacher_tag(sl, role="professor", day="Thu/Fri", slides_range=None):
    """
    GCP-style teacher tag.
    role : 'professor' | 'delegate'
    """
    col   = GCPBLUE if role == "professor" else GCPGREEN
    label = f"PROFESSOR — {day}"
    if slides_range:
        label += f" | Slides {slides_range}"
    if role == "delegate":
        label = f"DELEGATE — {day}"
        if slides_range:
            label += f" | Slides {slides_range}"
    tw = Inches(2.8)
    th = Pt(20)
    rect(sl, W - tw - Pt(8), Pt(4), tw, th, col, rounded=True, radius=0.06)
    txbox(sl, W - tw - Pt(8), Pt(4), tw, th,
          label, size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
