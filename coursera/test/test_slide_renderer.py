"""Regression tests for the trusted, data-only PowerPoint renderer."""

import builtins
from copy import deepcopy
import importlib
import json
from pathlib import Path

import pytest

from slide_gen import renderer


@pytest.fixture
def deck():
    return {
        "title": "A trusted deck",
        "slides": [{
            "title": "Introduction",
            "bullets": ["First point", "Second point"],
            "notes": "Speaker notes\nSecond line: <literal> & Unicode \u4e2d\u6587.",
        }],
    }


@pytest.fixture
def pptx():
    return pytest.importorskip("pptx")


@pytest.mark.parametrize("payload", [
    None, True, 12, [], (), "import os; os.remove('important')",
    "```python\nprint('not a deck')\n```", "{'title': 'Python'}",
    '{"title": "JSON source is not a decoded object", "slides": []}',
])
def test_reject_non_object_payloads(payload):
    with pytest.raises(ValueError):
        renderer.validate_deck(payload)


@pytest.mark.parametrize("key,value", [
    ("title", ""), ("title", " \n\t"), ("title", "x" * 121),
    ("title", 1), ("title", None), ("slides", []), ("slides", {}),
    ("slides", "Python"), ("slides", ()), ("slides", [None]),
])
def test_reject_invalid_root_fields(deck, key, value):
    deck[key] = value
    with pytest.raises(ValueError):
        renderer.validate_deck(deck)


@pytest.mark.parametrize("key,value", [
    ("title", "x" * 101), ("title", None), ("title", []),
    ("bullets", []), ("bullets", ["x"] * 7), ("bullets", "code"),
    ("bullets", ("x",)), ("bullets", [""]), ("bullets", [" \n"]),
    ("bullets", ["x" * 401]), ("bullets", [123]), ("bullets", [["x"]]),
    ("notes", "x" * 8001), ("notes", None), ("notes", {}),
])
def test_reject_invalid_slide_fields(deck, key, value):
    deck["slides"][0][key] = value
    with pytest.raises(ValueError):
        renderer.validate_deck(deck)


@pytest.mark.parametrize("location,key", [
    ("root", "title"), ("root", "slides"), ("slide", "title"),
    ("slide", "bullets"), ("slide", "notes"),
])
def test_all_schema_keys_required(deck, location, key):
    target = deck if location == "root" else deck["slides"][0]
    del target[key]
    with pytest.raises(ValueError):
        renderer.validate_deck(deck)


@pytest.mark.parametrize("location", ["root", "slide"])
@pytest.mark.parametrize("key", ["code", "python", "palette", "output_path", "__class__"])
def test_unknown_keys_rejected(deck, location, key):
    target = deck if location == "root" else deck["slides"][0]
    target[key] = "arbitrary code or configuration"
    with pytest.raises(ValueError):
        renderer.validate_deck(deck)


@pytest.mark.parametrize("text", ["\x00", "\x08", "\x0b", "\ud800", "\ufffe", "\uffff"])
@pytest.mark.parametrize("field", ["deck_title", "title", "bullet", "notes"])
def test_xml_invalid_characters_rejected(deck, text, field):
    if field == "deck_title":
        deck["title"] = text
    elif field == "bullet":
        deck["slides"][0]["bullets"] = [text]
    else:
        deck["slides"][0][field] = text
    with pytest.raises(ValueError):
        renderer.validate_deck(deck)


def test_schema_boundaries_and_detached_copy(deck):
    deck["title"] = "x" * 120
    deck["slides"][0] = {
        "title": "x" * 100, "bullets": ["x" * 400] * 6, "notes": "x" * 8000,
    }
    result = renderer.validate_deck(deck)
    assert result == deck
    result["slides"][0]["bullets"].append("changed")
    assert len(deck["slides"][0]["bullets"]) == 6
    deck["slides"] = [{"title": "", "bullets": ["x"], "notes": ""}] * 60
    assert len(renderer.validate_deck(deck)["slides"]) == 60
    deck["slides"].append(deck["slides"][0])
    with pytest.raises(ValueError):
        renderer.validate_deck(deck)


def test_total_text_limit():
    slide = {"title": "", "bullets": ["x"], "notes": "x" * 7999}
    data = {"title": "x", "slides": [deepcopy(slide) for _ in range(25)]}
    data["slides"][0]["notes"] = "x" * 7998
    assert renderer.validate_deck(data) == data
    data["title"] += "x"
    with pytest.raises(ValueError, match="200000"):
        renderer.validate_deck(data)


def test_import_and_validation_do_not_require_pptx(deck, monkeypatch):
    original_import = builtins.__import__

    def without_pptx(name, *args, **kwargs):
        if name == "pptx" or name.startswith("pptx.") or "slide_utils" in name:
            raise AssertionError("PPTX dependencies must be lazy")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", without_pptx)
    importlib.reload(renderer)
    assert renderer.validate_deck(deck) == deck


@pytest.mark.parametrize("palette", ["unknown", "gcp", "", None, [], 1])
def test_unknown_palette_rejected_before_io(deck, tmp_path, palette):
    with pytest.raises(ValueError, match="palette"):
        renderer.render_deck(deck, tmp_path / "deck.pptx", palette=palette)
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("notes", ["", "\tNotes\r\nNext\rEnd", "x" * 8000])
def test_saved_deck_reopens_with_notes(deck, tmp_path, pptx, notes):
    deck["slides"][0]["notes"] = notes
    original = deepcopy(deck)
    destination = tmp_path / "deck.pptx"
    destination.write_bytes(b"old destination")
    result = renderer.render_deck(deck, destination)
    assert result == str(destination.resolve())
    assert deck == original
    reopened = pptx.Presentation(result)
    assert len(reopened.slides) == 1
    assert reopened.core_properties.title == deck["title"]
    assert reopened.slides[0].has_notes_slide
    assert reopened.slides[0].notes_slide.notes_text_frame.text == notes.replace(
        "\r\n", "\n"
    ).replace("\r", "\n")
    assert list(tmp_path.iterdir()) == [destination]


def test_python_injection_is_literal_text(deck, tmp_path, pptx):
    marker = tmp_path / "executed.txt"
    payload = f"__import__('pathlib').Path({str(marker)!r}).touch()"
    deck["slides"][0]["bullets"] = [payload]
    deck["slides"][0]["notes"] = payload
    output = renderer.render_deck(deck, tmp_path / "literal.pptx")
    assert not marker.exists()
    slide = pptx.Presentation(output).slides[0]
    bullet = next(shape for shape in slide.shapes if shape.name == "Bullet 1")
    literal = bullet.text.removeprefix("\u2022 ").replace("\v  ", "")
    assert literal == payload
    assert slide.notes_slide.notes_text_frame.text == payload
    assert not any(relationship.is_external for relationship in slide.part.rels.values())


@pytest.mark.parametrize("palette,background,accent,foreground", [
    ("IBM New", "F2F4F8", "0F62FE", "081830"),
    ("GCP", "F1F3F4", "1A73E8", "212226"),
    ("IBM Original", "F0F2F8", "1A56DB", "0D1B4B"),
])
def test_palette_colors(deck, tmp_path, pptx, palette, background, accent, foreground):
    output = renderer.render_deck(deck, tmp_path / "colors.pptx", palette=palette)
    slide = pptx.Presentation(output).slides[0]
    assert str(slide.shapes[0].fill.fore_color.rgb) == background
    assert str(slide.shapes[1].fill.fore_color.rgb) == accent
    for shape in slide.shapes:
        if shape.name in ("Title", "Bullet 1", "Footer"):
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    assert str(run.font.color.rgb) == foreground
    from slide_gen import slide_utils

    assert str(slide_utils.LT_GCPRED) == "FC8E8E"


@pytest.mark.parametrize("bullet", [
    "W" * 400, "\u754c" * 400, "\U0001f600" * 400,
    ("word " * 80).strip(), "a\n\t" * 133,
], ids=["long-token", "cjk", "emoji", "words", "whitespace"])
def test_dense_layout_paginates_deterministically(deck, tmp_path, pptx, bullet):
    deck["title"] = "\u754c" * 120
    deck["slides"][0]["title"] = "\u754c" * 100
    deck["slides"][0]["bullets"] = [bullet] * 6
    first = pptx.Presentation(renderer.render_deck(deck, tmp_path / "first.pptx"))
    second = pptx.Presentation(renderer.render_deck(deck, tmp_path / "second.pptx"))
    assert len(first.slides) > 1
    assert [[shape.text for shape in slide.shapes] for slide in first.slides] == [
        [shape.text for shape in slide.shapes] for slide in second.slides
    ]
    rendered_bullets = []
    for index, slide in enumerate(first.slides):
        assert slide.notes_slide.notes_text_frame.text == deck["slides"][0]["notes"]
        text_shapes = []
        for shape in slide.shapes:
            assert 0 <= shape.left < first.slide_width
            assert 0 <= shape.top < first.slide_height
            assert shape.width > 0 and shape.height > 0
            assert shape.left + shape.width <= first.slide_width
            assert shape.top + shape.height <= first.slide_height
            if not shape.text:
                continue
            text_shapes.append(shape)
            assert shape.text_frame.auto_size == pptx.enum.text.MSO_AUTO_SIZE.NONE
            assert shape.text_frame.word_wrap is False
            lines = shape.text.split("\v")
            if shape.name.startswith("Bullet"):
                assert all(renderer._units(line) <= 88 for line in lines)
                assert shape.top.pt >= renderer.BODY_TOP
                assert shape.top.pt + shape.height.pt <= renderer.BODY_BOTTOM
                assert len(lines) * renderer.BODY_LINE_HEIGHT <= shape.height.pt
                rendered_bullets.append(" ".join(
                    [lines[0][2:]] + [line[2:] for line in lines[1:]]
                ))
            elif shape.name == "Title":
                assert all(renderer._units(line) <= 64 for line in lines)
                assert len(lines) * 28 <= shape.height.pt
                if index:
                    assert f"(continued {index + 1})" in " ".join(lines)
            elif shape.name == "Footer":
                assert all(renderer._units(line) <= 128 for line in lines)
                assert len(lines) * 12 <= shape.height.pt
        for position, shape in enumerate(text_shapes):
            for other in text_shapes[position + 1:]:
                assert (
                    shape.left + shape.width <= other.left
                    or other.left + other.width <= shape.left
                    or shape.top + shape.height <= other.top
                    or other.top + other.height <= shape.top
                )
    assert len(rendered_bullets) == 6
    assert all(text.replace(" ", "") == "".join(bullet.split()) for text in rendered_bullets)


@pytest.mark.parametrize("existing", [True, False])
@pytest.mark.parametrize("failure", [
    "save", "corrupt", "empty", "missing_notes", "changed_notes", "replace",
])
def test_failure_preserves_destination_and_cleans_temp(
    deck, tmp_path, pptx, monkeypatch, failure, existing,
):
    from pptx.presentation import Presentation
    from slide_gen import slide_utils

    destination = tmp_path / "deck.pptx"
    if existing:
        destination.write_bytes(b"previous deck bytes")
    original_save = Presentation.save

    def broken_save(presentation, path):
        Path(path).write_bytes(b"partial or corrupt PPTX")
        if failure == "save":
            raise OSError("simulated save failure")

    def empty_save(presentation, path):
        original_save(pptx.Presentation(), path)

    def fail_replace(source, target):
        raise OSError("simulated replace failure")

    if failure in ("save", "corrupt"):
        monkeypatch.setattr(Presentation, "save", broken_save)
    elif failure == "empty":
        monkeypatch.setattr(Presentation, "save", empty_save)
    elif failure == "missing_notes":
        monkeypatch.setattr(slide_utils, "set_notes", lambda slide, notes: None)
    elif failure == "changed_notes":
        original_set_notes = slide_utils.set_notes
        monkeypatch.setattr(
            slide_utils, "set_notes",
            lambda slide, notes: original_set_notes(slide, "wrong"),
        )
    else:
        monkeypatch.setattr(renderer.os, "replace", fail_replace)
    with pytest.raises((OSError, ValueError, pptx.exc.PackageNotFoundError)):
        renderer.render_deck(deck, destination)
    if existing:
        assert destination.read_bytes() == b"previous deck bytes"
    else:
        assert not destination.exists()
    assert list(tmp_path.iterdir()) == ([destination] if existing else [])


def test_unique_temp_in_output_parent_validated_before_replace(deck, tmp_path, pptx, monkeypatch):
    original_validate = renderer._validate_saved
    original_replace = renderer.os.replace
    validated_paths = []
    destination = tmp_path / "deck.pptx"

    def validate(path, expected_notes):
        assert path.parent == tmp_path
        assert path.suffix == ".pptx"
        assert path != destination
        original_validate(path, expected_notes)
        validated_paths.append(path)

    def replace(source, target):
        assert source == validated_paths[-1]
        assert target == destination
        original_replace(source, target)

    monkeypatch.setattr(renderer, "_validate_saved", validate)
    monkeypatch.setattr(renderer.os, "replace", replace)
    renderer.render_deck(deck, destination)
    renderer.render_deck(deck, destination)
    assert len(set(validated_paths)) == 2
    assert all(not path.exists() for path in validated_paths)


def test_invalid_input_never_touches_destination(deck, tmp_path):
    destination = tmp_path / "deck.pptx"
    destination.write_bytes(b"old")
    with pytest.raises(ValueError):
        renderer.render_deck(json.dumps(deck), destination)
    assert destination.read_bytes() == b"old"
    assert list(tmp_path.iterdir()) == [destination]


def test_adversarial_word_boundaries_fit_reserved_regions(deck, tmp_path, pptx):
    deck["title"] = ("\u754c" * 33 + " ") * 3
    deck["slides"][0]["title"] = ("\u754c" * 17 + " ") * 5
    bullet = ("\u754c" * 22 + " ") * 17
    deck["slides"][0]["bullets"] = [bullet]
    output = renderer.render_deck(deck, tmp_path / "word-boundaries.pptx")
    presentation = pptx.Presentation(output)
    assert len(presentation.slides) == 2
    contents = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            lines = shape.text.split("\v")
            if shape.name == "Title":
                assert len(lines) * 28 <= shape.height.pt
            elif shape.name == "Footer":
                assert len(lines) * 12 <= shape.height.pt
            elif shape.name.startswith("Bullet"):
                assert shape.top.pt + shape.height.pt <= renderer.BODY_BOTTOM
                assert len(lines) * renderer.BODY_LINE_HEIGHT <= shape.height.pt
                contents.extend(line[2:] for line in lines)
    assert "".join(contents).replace(" ", "") == bullet.replace(" ", "")


def test_source_slide_order_and_notes(deck, tmp_path, pptx, monkeypatch):
    deck["slides"][0]["bullets"] = ["\u754c" * 400] * 6
    deck["slides"].append({"title": "Last slide", "bullets": ["Last point"], "notes": "Last notes"})
    monkeypatch.chdir(tmp_path)
    output = renderer.render_deck(deck, "relative.pptx")
    assert output == str(tmp_path / "relative.pptx")
    presentation = pptx.Presentation(output)
    assert len(presentation.slides) == 7
    notes = [slide.notes_slide.notes_text_frame.text for slide in presentation.slides]
    assert notes == [deck["slides"][0]["notes"]] * 6 + ["Last notes"]
    assert any(shape.text == "Last slide" for shape in presentation.slides[-1].shapes)
